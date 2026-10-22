#!/usr/bin/env python3
r"""Build the live-workbook snapshot used by the React dashboard.

The shared planning workbook is an Excel matrix of monthly targets and
accomplishments. This script reads the consolidated "Main (for CO)" sheet,
normalizes the merged/forward-filled labels, and writes a JSON snapshot that
the browser can render without needing Excel or a tenant token.

Usage:
    python scripts/sync_live_data.py
    python scripts/sync_live_data.py --source "C:\path\to\workbook.xlsx"
    python scripts/sync_live_data.py --source workbook.xlsx --out public/live-data.json

Requires: openpyxl (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from live_projects import build_projects

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - dependency guidance
    raise SystemExit(
        "openpyxl is required. Install it with: python -m pip install openpyxl"
    ) from exc


WORKBOOK_NAME = "EMB 7 Accomplishment 2026 - ALL SECTIONS.xlsx"
DEFAULT_OUTPUT = Path("public/live-data.json")
MONTHS_SHORT = [
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
]
MONTHS_LONG = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]
PLACEHOLDERS = {"", "-", "–", "—", "n/a", "na", "none", "null"}


def candidate_sources(cli_source: str | None) -> list[Path]:
    candidates: list[Path] = []
    if cli_source:
        candidates.append(Path(cli_source))
    env_source = os.environ.get("LIVE_WORKBOOK_PATH")
    if env_source:
        candidates.append(Path(env_source))

    candidates.append(Path("data") / WORKBOOK_NAME)
    home = Path.home()
    candidates.extend(
        [
            home
            / "OneDrive - Environmental Management Bureau"
            / "EMB Office Files"
            / "10 Reports (Annual, Monthly, Quarterly)"
            / "2026 Monthly Accomplishment"
            / "ORD-PLANNING"
            / WORKBOOK_NAME,
            home
            / "OneDrive - Environmental Management Bureau"
            / "EMB Office Files"
            / "10 Reports (Annual, Monthly, Quarterly)"
            / "2026 Monthly Accomplishment"
            / WORKBOOK_NAME,
            home
            / "OneDrive - Environmental Management Bureau"
            / WORKBOOK_NAME,
        ]
    )
    return candidates


def find_source(cli_source: str | None) -> Path:
    for candidate in candidate_sources(cli_source):
        if candidate.is_file():
            return candidate.resolve()
    looked = "\n".join(f"  - {p}" for p in candidate_sources(cli_source))
    raise SystemExit(
        "Workbook not found. Pass --source or set LIVE_WORKBOOK_PATH.\n"
        f"Looked in:\n{looked}"
    )


def text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def number(value: Any) -> float:
    if value is None or isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip()
    if raw.lower() in PLACEHOLDERS:
        return 0.0
    cleaned = (
        raw.replace(",", "")
        .replace("₱", "")
        .replace("%", "")
        .replace(" ", "")
    )
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def meaningful(value: Any) -> str:
    value_text = text(value)
    if value_text.lower() in PLACEHOLDERS:
        return ""
    if re.fullmatch(r"0(?:\.0+)?", value_text):
        return ""
    return value_text


def normalize_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def read_focals(workbook) -> dict[str, str]:
    """Return a normalized section key -> focal staff string.

    The Focals sheet uses slightly different labels from the Main sheet
    (e.g. "EMED-AWMS" vs "EMED - AWMS"). Normalization handles punctuation
    and spacing. AMBIENT and FAD have sub-teams; combine their focal names.
    """
    if "Focals" not in workbook.sheetnames:
        return {}
    ws = workbook["Focals"]
    leads: dict[str, list[str]] = defaultdict(list)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=7, values_only=True):
        section = text(row[1]) if len(row) > 1 else ""
        focal = meaningful(row[2]) if len(row) > 2 else ""
        if not section or not focal:
            continue
        key = normalize_key(section)
        # Roll sub-teams up to the section names used by the Main sheet.
        if key.startswith("fad"):
            key = "fad"
        if key.startswith("ambient"):
            key = "ambient"
        if key not in leads or focal not in leads[key]:
            leads[key].append(focal)

    return {key: " / ".join(values) for key, values in leads.items()}


def build_payload(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if "Main (for CO)" not in workbook.sheetnames:
        raise SystemExit(
            'The workbook has no "Main (for CO)" sheet. '
            f"Found: {', '.join(workbook.sheetnames)}"
        )

    ws = workbook["Main (for CO)"]
    focals = read_focals(workbook)
    rows = list(
        ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=48, values_only=True)
    )
    if len(rows) < 5:
        raise SystemExit("The Main sheet does not contain the expected header/data rows.")

    section = ""
    project = ""
    raw_records: list[dict[str, Any]] = []
    month_target = [0.0] * 12
    month_actual = [0.0] * 12

    # Data starts on row 5 in the current EMB template. Rows 1-4 are headers.
    for excel_row, row in enumerate(rows[4:], start=5):
        a = text(row[0]) if len(row) > 0 else ""
        b = text(row[1]) if len(row) > 1 else ""
        c = text(row[2]) if len(row) > 2 else ""

        if a:
            if not section or normalize_key(a) != normalize_key(section):
                project = ""
            section = a
        if b:
            project = b

        targets = [number(row[3 + i]) if len(row) > 3 + i else 0.0 for i in range(12)]
        actuals = [number(row[17 + i]) if len(row) > 17 + i else 0.0 for i in range(12)]

        if not c:
            continue

        annual_target = round(sum(targets), 2)
        annual_actual = round(sum(actuals), 2)

        remarks_a = meaningful(row[32]) if len(row) > 32 else ""
        mov = meaningful(row[33]) if len(row) > 33 else ""
        remarks_b = meaningful(row[34]) if len(row) > 34 else ""
        remarks = remarks_b or remarks_a
        if remarks_a and remarks_b and remarks_a != remarks_b:
            remarks = f"{remarks_b} / {remarks_a}"

        for i in range(12):
            month_target[i] += targets[i]
            month_actual[i] += actuals[i]

        raw_records.append(
            {
                "row": excel_row,
                "section": section,
                "project": project,
                "indicator": c,
                "targets": [round(v, 2) for v in targets],
                "actuals": [round(v, 2) for v in actuals],
                "annualTarget": annual_target,
                "annualActual": annual_actual,
                "remarks": remarks,
                "mov": mov,
            }
        )

    # The latest month with an actual value is the reporting month. October to
    # December are normally empty when the snapshot is taken during the year.
    reported_months = [i for i, value in enumerate(month_actual) if abs(value) > 0.0001]
    reporting_index = reported_months[-1] if reported_months else 0
    reporting_month_number = reporting_index + 1
    reporting_year = 2026
    reporting_label = f"{MONTHS_LONG[reporting_index]} {reporting_year}"

    # Monthly medians are safer than monthly sums because the indicators use
    # different units (people reached, reports, inspections, hectares, etc.).
    monthly: list[dict[str, Any]] = []
    for month_index in range(12):
        rates: list[float] = []
        met = 0
        target_lines = 0
        for raw in raw_records:
            target = raw["targets"][month_index]
            actual = raw["actuals"][month_index]
            if target > 0.0001:
                target_lines += 1
                rates.append(actual / target)
                if actual >= target:
                    met += 1
        monthly.append(
            {
                "index": month_index + 1,
                "label": MONTHS_SHORT[month_index],
                "assessed": len(rates),
                "met": met,
                "metRate": round(met / len(rates), 4) if rates else None,
                "medianRate": round(statistics.median(rates), 4) if rates else None,
                "isReportingMonth": month_index == reporting_index,
                "isFuture": month_index > reporting_index,
            }
        )

    records: list[dict[str, Any]] = []
    for raw in raw_records:
        ytd_target = round(sum(raw["targets"][: reporting_index + 1]), 2)
        ytd_actual = round(sum(raw["actuals"][: reporting_index + 1]), 2)
        annual_target = raw["annualTarget"]
        annual_actual = raw["annualActual"]

        ytd_rate = ytd_actual / ytd_target if ytd_target > 0.0001 else None
        annual_rate = annual_actual / annual_target if annual_target > 0.0001 else None

        if annual_target <= 0.0001 and annual_actual > 0.0001:
            status = "Unplanned"
        elif annual_target <= 0.0001:
            status = "No Target"
        elif annual_actual >= annual_target:
            status = "Completed"
        elif ytd_target <= 0.0001:
            status = "Not Started"
        elif ytd_rate is not None and ytd_rate >= 0.9:
            status = "On Track"
        elif ytd_rate is not None and ytd_rate >= 0.7:
            status = "At Risk"
        else:
            status = "Delayed"

        reported = [i for i, value in enumerate(raw["actuals"]) if value > 0.0001]
        latest_actual = reported[-1] if reported else None
        latest_label = (
            f"{MONTHS_SHORT[latest_actual]} {reporting_year}"
            if latest_actual is not None
            else "No actual"
        )

        flags: list[str] = []
        if annual_target <= 0.0001:
            flags.append("No target")
        if annual_target > 0.0001 and annual_actual <= 0.0001:
            flags.append("No accomplishment")
        if not raw["remarks"]:
            flags.append("Missing remarks")
        if not raw["mov"]:
            flags.append("Missing MOV")

        section_key = normalize_key(raw["section"])
        lead = focals.get(section_key, "")
        if not lead and section_key.startswith("cpe"):
            lead = focals.get("cpdeia", "")
        if not lead and section_key == "cpdeia":
            lead = focals.get("eia", "")

        records.append(
            {
                "id": f"ACT-2026-{raw['row']:03d}",
                "row": raw["row"],
                "section": raw["section"],
                "project": raw["project"],
                "indicator": raw["indicator"],
                "lead": lead,
                "targets": raw["targets"],
                "actuals": raw["actuals"],
                "annualTarget": annual_target,
                "annualActual": annual_actual,
                "ytdTarget": ytd_target,
                "ytdActual": ytd_actual,
                "ytdRate": None if ytd_rate is None else round(ytd_rate, 4),
                "annualRate": None if annual_rate is None else round(annual_rate, 4),
                "status": status,
                "latestActualMonth": latest_label,
                "remarks": raw["remarks"],
                "mov": raw["mov"],
                "flags": flags,
            }
        )

    # Section summaries use the median of assessed indicators so a single
    # high-volume activity (for example social-media reach) does not dominate.
    by_section: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_section[record["section"]].append(record)

    sections: list[dict[str, Any]] = []
    for name, items in by_section.items():
        assessed = [
            item["ytdRate"]
            for item in items
            if item["ytdRate"] is not None and item["status"] not in {"No Target"}
        ]
        rates = [float(rate) for rate in assessed if rate is not None]
        status_counts = Counter(item["status"] for item in items)
        lead = next((item["lead"] for item in items if item["lead"]), "")
        sections.append(
            {
                "section": name,
                "activities": len(items),
                "medianRate": round(statistics.median(rates), 4) if rates else None,
                "onTrack": status_counts.get("On Track", 0),
                "atRisk": status_counts.get("At Risk", 0),
                "delayed": status_counts.get("Delayed", 0),
                "completed": status_counts.get("Completed", 0),
                "notStarted": status_counts.get("Not Started", 0),
                "unplanned": status_counts.get("Unplanned", 0),
                "noTarget": status_counts.get("No Target", 0),
                "lead": lead,
            }
        )
    sections.sort(
        key=lambda item: (
            item["medianRate"] is None,
            -(item["medianRate"] or 0),
            item["section"].lower(),
        )
    )

    all_rates = [
        record["ytdRate"]
        for record in records
        if record["ytdRate"] is not None and record["status"] != "No Target"
    ]
    all_rates_float = [float(rate) for rate in all_rates]
    status_counts = Counter(record["status"] for record in records)
    data_quality = {
        "missingRemarks": sum("Missing remarks" in record["flags"] for record in records),
        "missingMov": sum("Missing MOV" in record["flags"] for record in records),
        "noTarget": sum("No target" in record["flags"] for record in records),
        "noAccomplishment": sum(
            "No accomplishment" in record["flags"] for record in records
        ),
    }

    modified = datetime.fromtimestamp(source.stat().st_mtime).astimezone()
    generated = datetime.now().astimezone()

    # The real planning portfolio: the workbook's own activity groups.
    projects = build_projects(records, reporting_index, reporting_year)
    units: list[str] = []
    for record in records:
        if record["section"] and record["section"] not in units:
            units.append(record["section"])

    return {
        "meta": {
            "sourceFile": source.name,
            "sourceFolder": source.parent.name,
            "sourceSheet": "Main (for CO)",
            "reportingYear": reporting_year,
            "reportingMonth": reporting_month_number,
            "reportingIndex": reporting_index,
            "reportingLabel": reporting_label,
            "monthsReported": reporting_index + 1,
            "monthLabels": MONTHS_SHORT,
            "monthLabelsLong": [
                f"{month} {reporting_year}" for month in MONTHS_LONG
            ],
            "generatedAt": generated.isoformat(timespec="seconds"),
            "workbookModifiedAt": modified.isoformat(timespec="seconds"),
            "recordCount": len(records),
            "sectionCount": len(sections),
            "projectCount": len(projects),
            "unitCount": len(units),
            "medianYtdRate": (
                round(statistics.median(all_rates_float), 4)
                if all_rates_float
                else None
            ),
            "averageYtdRate": (
                round(statistics.fmean(all_rates_float), 4)
                if all_rates_float
                else None
            ),
            "statusCounts": dict(status_counts),
            "dataQuality": data_quality,
            "focalNote": "Focal staff from the workbook's Focals sheet (FY 2025).",
            "snapshotNote": (
                "Snapshot generated from the local OneDrive copy of the workbook. "
                "Run `npm run sync:live` to refresh it."
            ),
        },
        "monthly": monthly,
        "sections": sections,
        "records": records,
        "projects": projects,
        "units": units,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", help="Path to the live .xlsx workbook")
    parser.add_argument(
        "--out",
        default=str(DEFAULT_OUTPUT),
        help=f"Output JSON path (default: {DEFAULT_OUTPUT})",
    )
    args = parser.parse_args()

    source = find_source(args.source)
    payload = build_payload(source)
    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    meta = payload["meta"]
    print(f"Wrote {output.resolve()}")
    print(
        f"{meta['recordCount']} records, {meta['sectionCount']} sections, "
        f"reporting through {meta['reportingLabel']}"
    )
    print(f"Source: {source}")


if __name__ == "__main__":
    main()
